from datetime import date
from decimal import Decimal
import json
from pathlib import Path
import tempfile
import unittest

import openpyxl

import market_report as report


class ReportingTests(unittest.TestCase):
    def setUp(self):
        self.settings = json.loads((report.BASE / "settings.json").read_text())

    def row(self, mls="1", status="CLSD", community="Carillon", city="Plainfield", closed="2026-09-30"):
        return dict(mls=mls, status=status, community=community, city=city,
                    closed_date=closed, list_price=200000, original_list_price=200000, sold_price=100000, dom=0,
                    street_number=123, street_name="Main", street_suffix="St", unit="B",
                    location="fixture row", source="fixture.xlsx", row=2)

    def test_aggregate_is_not_average_of_percentages(self):
        a, b = self.row(), self.row("2")
        b.update(list_price=300000, original_list_price=300000, sold_price=300000, dom=10)
        buckets, _ = report.classify({"closed": [a, b]}, self.settings, date(2026, 10, 1))
        result = report.summarize(self.settings["communities"][0], buckets["carillon"], date(2026, 10, 1))
        self.assertEqual(result["aggregate_sale_to_list_percentage"], 80)
        self.assertEqual(result["median_closed_dom"], 5)
        self.assertEqual(result["median_sold_price"], 200000)

    def test_date_boundaries(self):
        dates = ["2025-09-30", "2025-10-01", "2026-09-30", "2026-10-01"]
        rows = [self.row(str(i), closed=d) for i, d in enumerate(dates)]
        buckets, review = report.classify({"closed": rows}, self.settings, date(2026, 10, 1))
        self.assertEqual(len(buckets["carillon"]["closed"]), 2)
        self.assertEqual(sum(r["disposition"] != "included" for r in review), 2)
        self.assertEqual(report.year_before(date(2024, 2, 29)), date(2023, 2, 28))

    def test_exact_city_alias_and_status_groups(self):
        rows = [self.row(str(i), status=s) for i, s in enumerate(self.settings["statuses"]["pending"])]
        rows += [self.row("a", "NEW", " CARILLON  "), self.row("b", "PCHG"),
                 self.row("c", "ACTV", "Carillon Club", "Naperville"),
                 self.row("d", "ACTV", "Carillon", "Aurora"),
                 self.row("e", "ACTV", "Shorewood Glen Del Webb", "Shorewood")]
        buckets, review = report.classify({"inventory": rows}, self.settings, date(2026, 10, 1))
        self.assertEqual(len(buckets["carillon"]["pending"]), 7)
        self.assertEqual(len(buckets["carillon"]["active"]), 2)
        self.assertEqual(len(buckets["carillon-club"]["active"]), 1)
        self.assertEqual(len(buckets["shorewood-glen"]["active"]), 1)
        self.assertEqual(sum(r["disposition"] != "included" for r in review), 1)

    def test_duplicate_unknown_status_and_missing_price_stop(self):
        for rows in ([self.row(), self.row()], [self.row(status="UNKNOWN")],
                     [dict(self.row(), list_price=None)], [dict(self.row(), sold_price=0)],
                     [dict(self.row(), dom=-1)], [dict(self.row(), closed_date=None)],
                     [dict(self.row(), community=None)]):
            with self.subTest(rows=rows), self.assertRaises(report.ReportError):
                report.classify({"closed": rows}, self.settings, date(2026, 10, 1))

    def test_empty_community(self):
        result = report.summarize(self.settings["communities"][0],
                                  dict(active=[], pending=[], closed=[]), date(2026, 10, 1))
        self.assertEqual(result["active_listings"], 0)
        self.assertIsNone(result["median_sold_price"])
        self.assertIsNone(result["aggregate_sale_to_list_percentage"])

    def test_latest_month_calendar_boundaries(self):
        for as_of, start, end in [
            (date(2026, 10, 1), date(2026, 9, 1), date(2026, 10, 1)),
            (date(2026, 10, 15), date(2026, 9, 1), date(2026, 10, 1)),
            (date(2027, 1, 1), date(2026, 12, 1), date(2027, 1, 1)),
            (date(2024, 3, 1), date(2024, 2, 1), date(2024, 3, 1)),
        ]:
            with self.subTest(as_of=as_of):
                self.assertEqual(report.latest_month_bounds(as_of), (start, end))

    def test_latest_month_aggregate_and_inclusive_boundaries(self):
        rows = [dict(self.row('a', closed='2026-09-01'), sold_price=100000, original_list_price=200000, dom=10),
                dict(self.row('b', closed='2026-09-30'), sold_price=300000, original_list_price=300000, dom=30),
                dict(self.row('c', closed='2026-08-31'), sold_price=900000, original_list_price=1000000, dom=90),
                dict(self.row('d', closed='2026-10-01'), sold_price=900000, original_list_price=1000000, dom=90)]
        buckets, _ = report.classify({'closed': rows}, self.settings, date(2026, 10, 15))
        result = report.latest_month_statistics(buckets['carillon']['closed'], date(2026, 10, 15))
        self.assertEqual(result['latest_complete_month'], '2026-09')
        self.assertEqual(result['latest_month_closed_sales'], 2)
        self.assertEqual(result['latest_month_median_sold_price'], 200000)
        self.assertEqual(result['latest_month_sale_to_list_percentage'], 80)
        self.assertEqual(result['latest_month_median_closed_dom'], 20)

    def test_latest_month_zero_does_not_carry_forward(self):
        buckets, _ = report.classify({'closed': [self.row(closed='2026-08-31')]}, self.settings, date(2026, 10, 1))
        result = report.summarize(self.settings['communities'][0], buckets['carillon'], date(2026, 10, 1))
        self.assertEqual(result['closed_sales_12_months'], 1)
        self.assertEqual(result['median_sold_price'], 100000)
        self.assertEqual(result['latest_month_closed_sales'], 0)
        for field in ('latest_month_median_sold_price', 'latest_month_sale_to_list_percentage', 'latest_month_median_closed_dom'):
            self.assertIsNone(result[field])

    def test_latest_month_leap_day_is_included(self):
        buckets, _ = report.classify({'closed': [self.row(closed='2024-02-29')]}, self.settings, date(2024, 3, 1))
        result = report.summarize(self.settings['communities'][0], buckets['carillon'], date(2024, 3, 1))
        self.assertEqual(result['latest_complete_month'], '2024-02')
        self.assertEqual(result['latest_month_closed_sales'], 1)

    def test_latest_month_invalid_original_price_stops_run(self):
        row = dict(self.row(closed='2026-09-15'), original_list_price=None)
        with self.assertRaises(report.OriginalPriceError):
            report.classify({'closed': [row]}, self.settings, date(2026, 10, 1))

    def test_original_price_not_regular_list_price(self):
        row = self.row()
        row.update(original_list_price=250000, list_price=200000, sold_price=200000)
        buckets, _ = report.classify({"closed": [row]}, self.settings, date(2026, 10, 1))
        result = report.summarize(self.settings["communities"][0], buckets["carillon"], date(2026, 10, 1))
        self.assertEqual(result["aggregate_sale_to_list_percentage"], 80)
        self.assertEqual(result["total_closed_original_list_price"], 250000)

    def test_all_bad_original_prices_reported_without_fallback(self):
        bad = [None, "", 0, -5, "abc", "NaN", "Infinity", True]
        rows = [dict(self.row(str(i)), original_list_price=value) for i, value in enumerate(bad)]
        with self.assertRaises(report.ReportError) as raised:
            report.classify({"closed": rows}, self.settings, date(2026, 10, 1))
        for i in range(len(bad)):
            self.assertIn(f"MLS {i},", str(raised.exception))
        self.assertIn("8 qualifying", str(raised.exception))
        self.assertIn("123 Main St Unit B, Plainfield", str(raised.exception))
        self.assertIn("WARNING", str(raised.exception))

    def test_denominator_mapping_cannot_be_changed_in_settings(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self.settings['closed_columns']['original_list_price'] = 'List Price'
            config = root / 'settings.json'
            config.write_text(json.dumps(self.settings))
            with self.assertRaisesRegex(report.ReportError, 'Approved methodology requires Orig List Pr'):
                report.run(root / 'Input', root / 'Output', config, date(2026, 10, 1))

    def test_failed_validation_saves_warning_with_mls_and_address(self):
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            inputs = root / 'Input'
            inputs.mkdir()
            for name in self.settings['input_files'].values():
                (inputs / name).write_bytes(b'fixture')
            def fake_load(path, columns):
                rows = [dict(self.row('BAD123'), original_list_price=None)] if 'Closed' in path.name else []
                return rows, {'file': path.name}
            with patch.object(report, 'load_rows', side_effect=fake_load):
                with self.assertRaises(report.OriginalPriceError):
                    report.run(inputs, root / 'Output', report.BASE / 'settings.json', date(2026, 10, 1))
            files = list((root / 'Archive' / '_validation_warnings').iterdir())
            self.assertEqual(len(files), 1)
            self.assertIn('WARNING', files[0].name)
            warning = files[0].read_text()
            self.assertIn('BAD123', warning)
            self.assertIn('123 Main St Unit B, Plainfield', warning)

    def test_original_price_not_required_for_excluded_sales_or_inventory(self):
        rows = [dict(self.row("1", closed="2025-09-30"), original_list_price=None),
                dict(self.row("2", community="Other"), original_list_price=None)]
        active = dict(self.row("3", status="ACTV"), original_list_price=None)
        buckets, _ = report.classify({"closed": rows, "inventory": [active]}, self.settings, date(2026, 10, 1))
        self.assertEqual(len(buckets["carillon"]["active"]), 1)
        self.assertEqual(len(buckets["carillon"]["closed"]), 0)

    def test_numeric_validation(self):
        self.assertEqual(report.number("$250,000", "price", "test", True), Decimal(250000))
        for value in ("NaN", "Infinity", None, "not a number", True):
            with self.subTest(value=value), self.assertRaises(report.ReportError):
                report.number(value, "DOM", "test")

    def test_reader_accepts_new_sheet_name_and_rejects_changed_column(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "New MRED Sheet"
            ws.append(["Export heading"])
            ws.append(list(self.settings["columns"].values()))
            ws.append(["Carillon", "123", "ACTV", 0, None, 250000, None, 1, "Main", "St", None, "Plainfield"])
            wb.save(path)
            rows, info = report.load_rows(path, self.settings["columns"])
            self.assertEqual(len(rows), 1)
            self.assertEqual(info["header_row"], 2)
            ws.cell(2, 4, "Different DOM")
            wb.save(path)
            wb.close()
            with self.assertRaises(report.ReportError):
                report.load_rows(path, self.settings["columns"])

    def make_exports(self, root):
        inputs = root / 'Input'
        inputs.mkdir()
        for role, name in self.settings['input_files'].items():
            columns = dict(self.settings['columns'])
            if role == 'closed':
                columns.update(self.settings['closed_columns'])
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(list(columns.values()))
            row = self.row('sold' if role == 'closed' else 'active',
                           status='CLSD' if role == 'closed' else 'ACTV')
            ws.append([row.get(key) for key in columns])
            wb.save(inputs / name)
            wb.close()
        return inputs

    def test_monthly_archive_and_duplicate_date_protection(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            inputs = self.make_exports(root)
            output = root / 'Output'
            config = report.BASE / 'settings.json'
            report.run(inputs, output, config, date(2026, 10, 1))
            archive = root / 'Archive' / '2026-10-01'
            history = root / 'Archive' / 'market_history.xlsx'
            wb = openpyxl.load_workbook(history)
            original_rows = [(c.coordinate, c.value, c.number_format, str(c._style))
                             for row in wb['Market History'].iter_rows(min_row=2) for c in row]
            self.assertEqual(wb['Market History'].max_row, 1 + len(self.settings['communities']))
            wb.close()
            for path in output.rglob('*'):
                if path.is_file():
                    self.assertEqual(path.read_bytes(), (archive / path.relative_to(output)).read_bytes())
            for path in inputs.iterdir():
                self.assertEqual(path.read_bytes(), (archive / 'source_exports' / path.name).read_bytes())
            self.assertEqual(len(list((output / 'communities').glob('*.json'))), len(self.settings['communities']))
            for path in output.rglob('*.json'):
                self.assertEqual(json.loads(path.read_text())['report_date'], '2026-10-01')
            before = {p.relative_to(root): p.read_bytes() for p in root.rglob('*') if p.is_file()}
            with self.assertRaisesRegex(report.ReportError, 'WARNING: Archive already exists'):
                report.run(inputs, output, config, date(2026, 10, 1))
            self.assertEqual(before, {p.relative_to(root): p.read_bytes() for p in root.rglob('*') if p.is_file()})
            report.run(inputs, output, config, date(2026, 11, 1))
            wb = openpyxl.load_workbook(history)
            self.assertEqual(wb['Market History'].max_row, 1 + 2 * len(self.settings['communities']))
            self.assertEqual(original_rows, [(c.coordinate, c.value, c.number_format, str(c._style))
                                            for row in wb['Market History'].iter_rows(min_row=2, max_row=1 + len(self.settings['communities'])) for c in row])
            self.assertEqual(sum(row[1] == 'Lago Vista' for row in wb['Market History'].iter_rows(min_row=2, values_only=True)), 2)
            wb.close()
            self.assertEqual(json.loads((output / 'run_info.json').read_text())['report_date'], '2026-11-01')
            self.assertEqual(json.loads((archive / 'run_info.json').read_text())['report_date'], '2026-10-01')
            self.assertTrue((root / 'Archive' / '2026-11-01' / 'market_summary.xlsx').exists())
            self.assertEqual(set(p.name for p in output.iterdir()),
                             {'communities', 'audit', 'market_summary.xlsx', 'listing_review.csv', 'run_info.json'})

    def test_public_json_exact_allowlist_and_private_audit_preserved(self):
        expected = {'community', 'city', 'state', 'report_date', 'closed_period_start', 'closed_period_end',
                    'active_listings', 'median_active_list_price', 'average_active_dom',
                    'pending_contingent_fin_listings', 'closed_sales_12_months', 'median_sold_price',
                    'aggregate_sale_to_list_percentage', 'median_closed_dom', 'sale_to_list_basis',
                    'latest_complete_month', 'latest_month_closed_sales', 'latest_month_median_sold_price',
                    'latest_month_sale_to_list_percentage', 'latest_month_median_closed_dom', 'monthly_history', 'inventory_updated', 'closed_sales_through'}
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            reports = []
            for community in self.settings['communities']:
                full = report.summarize(community, dict(active=[], pending=[], closed=[]), date(2026, 10, 1))
                full['new_private_field'] = 'must never appear publicly'
                reports.append(full)
            report.write_community_reports(root, self.settings['communities'], reports)
            for community, full in zip(self.settings['communities'], reports):
                name = community['slug'] + '.json'
                public = json.loads((root / 'communities' / name).read_text())
                audit = json.loads((root / 'audit' / 'communities' / name).read_text())
                self.assertEqual(set(public), expected)
                self.assertEqual(len(public), 23)
                self.assertEqual(public, {key: full[key] for key in expected})
                self.assertEqual(audit, full)

    def test_publication_failure_restores_current_report(self):
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output, stage, archive_stage = root / 'Output', root / 'stage', root / 'archive-stage'
            archive_root = root / 'Archive'
            for path in (output, stage, archive_stage, archive_root):
                path.mkdir()
            (output / 'old.txt').write_text('preserve me')
            original_rename = Path.rename
            def failing_rename(path, target):
                if path == stage:
                    raise PermissionError('Simulated publish failure')
                return original_rename(path, target)
            with patch.object(Path, 'rename', failing_rename):
                with self.assertRaises(PermissionError):
                    report.publish_reports(stage, archive_stage, output, archive_root / '2026-10-01')
            self.assertEqual((output / 'old.txt').read_text(), 'preserve me')
            self.assertFalse((archive_root / '2026-10-01').exists())

    def test_nested_report_paths_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with self.assertRaisesRegex(report.ReportError, 'separate folders'):
                report.run(root / 'Input', root, report.BASE / 'settings.json', date(2026, 10, 1))

    def test_lago_vista_exact_match_and_all_metrics(self):
        rows = [dict(self.row('a', 'ACTV', ' Lago   Vista ', 'LOCKPORT'), list_price=300000, dom=10),
                dict(self.row('b', 'NEW', 'Lago Vista', 'Lockport'), list_price=400000, dom=20),
                self.row('c', 'FIN', 'Lago Vista', 'Lockport'),
                self.row('d', 'PEND', 'Lago Vista', 'Lockport'),
                self.row('e', 'ACTV', 'Lago Vista', 'Other'),
                self.row('f', 'ACTV', 'Lago Vista North', 'Lockport')]
        closed = [dict(self.row('s1', community='Lago Vista', city='Lockport'),
                       original_list_price=250000, sold_price=200000, dom=10),
                  dict(self.row('s2', community='Lago Vista', city='Lockport'),
                       original_list_price=350000, sold_price=300000, dom=30)]
        buckets, review = report.classify({'inventory': rows, 'closed': closed}, self.settings, date(2026, 10, 1))
        community = next(c for c in self.settings['communities'] if c['slug'] == 'lago-vista')
        result = report.summarize(community, buckets['lago-vista'], date(2026, 10, 1))
        expected = dict(active_listings=2, median_active_list_price=350000, average_active_dom=15,
                        pending_contingent_fin_listings=2, closed_sales_12_months=2,
                        median_sold_price=250000, aggregate_sale_to_list_percentage=83.3333, median_closed_dom=20)
        for key, value in expected.items():
            self.assertEqual(result[key], value)
        self.assertEqual(sum(r['disposition'] != 'included' for r in review), 2)
        closed[0]['original_list_price'] = None
        with self.assertRaisesRegex(report.OriginalPriceError, 'MLS s1, Lago Vista'):
            report.classify({'closed': closed}, self.settings, date(2026, 10, 1))

    def test_archive_history_import_rejects_existing_date(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            inputs = self.make_exports(root)
            report.run(inputs, root / 'Output', report.BASE / 'settings.json', date(2026, 10, 1))
            history = root / 'Archive' / 'market_history.xlsx'
            before = history.read_bytes()
            with self.assertRaisesRegex(report.ReportError, 'already exists in market_history'):
                report.import_archive_history(root / 'Archive' / '2026-10-01')
            self.assertEqual(before, history.read_bytes())

    def test_history_failure_rolls_back_report_publication(self):
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            inputs = self.make_exports(root)
            output = root / 'Output'
            report.run(inputs, output, report.BASE / 'settings.json', date(2026, 10, 1))
            history = root / 'Archive' / 'market_history.xlsx'
            before = history.read_bytes()
            current = (output / 'run_info.json').read_bytes()
            with patch.object(report, 'commit_history', side_effect=PermissionError('History open in Excel')):
                with self.assertRaises(PermissionError):
                    report.run(inputs, output, report.BASE / 'settings.json', date(2026, 11, 1))
            self.assertEqual(before, history.read_bytes())
            self.assertEqual(current, (output / 'run_info.json').read_bytes())
            self.assertFalse((root / 'Archive' / '2026-11-01').exists())

    def test_history_date_blocks_run_even_without_dated_archive(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            inputs = self.make_exports(root)
            output = root / 'Output'
            report.run(inputs, output, report.BASE / 'settings.json', date(2026, 10, 1))
            archive = root / 'Archive'
            (archive / '2026-10-01').rename(archive / 'preserved-archive')
            before = (archive / 'market_history.xlsx').read_bytes()
            with self.assertRaisesRegex(report.ReportError, 'already exists in market_history'):
                report.run(inputs, output, report.BASE / 'settings.json', date(2026, 10, 1))
            self.assertEqual(before, (archive / 'market_history.xlsx').read_bytes())
            self.assertFalse((archive / '2026-10-01').exists())


if __name__ == "__main__":
    unittest.main()
