"""Read MRED exports without changing them and create community market reports."""
from __future__ import annotations

import argparse
import calendar
from collections import Counter
import csv
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
import hashlib
import json
import math
from pathlib import Path
import statistics
import shutil
import sys
import tempfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = Path(__file__).resolve().parent
SALE_TO_LIST_DEFINITION = "Total Sold Price divided by Total Original List Price for qualifying closed sales."
# Approved methodology: change only on the user's explicit instruction.
ORIGINAL_LIST_FIELD = "Orig List Pr"
PUBLIC_COMMUNITY_FIELDS = (
    "community", "city", "state", "report_date", "closed_period_start", "closed_period_end",
    "active_listings", "median_active_list_price", "average_active_dom",
    "pending_contingent_fin_listings", "closed_sales_12_months", "median_sold_price",
    "aggregate_sale_to_list_percentage", "median_closed_dom", "sale_to_list_basis",
    "latest_complete_month", "latest_month_closed_sales", "latest_month_median_sold_price",
    "latest_month_sale_to_list_percentage", "latest_month_median_closed_dom",
    "monthly_history",
    "inventory_updated", "closed_sales_through",
)

MONTHLY_HISTORY_FIELDS = (
    "month", "closed_sales", "median_sold_price",
    "sale_to_original_list_percentage", "median_closed_dom",
)


def public_community_report(report):
    """Explicit allowlist: new internal fields can never leak into public JSON."""
    validate_monthly_history(report)
    validate_source_dates(report)
    return {field: report[field] for field in PUBLIC_COMMUNITY_FIELDS}


def write_community_reports(folder, communities, reports):
    public_dir = folder / 'communities'
    audit_dir = folder / 'audit' / 'communities'
    public_dir.mkdir(parents=True, exist_ok=True)
    audit_dir.mkdir(parents=True, exist_ok=True)
    for community, report in zip(communities, reports):
        filename = community['slug'] + '.json'
        (audit_dir / filename).write_text(json.dumps(report, indent=2, allow_nan=False) + '\n', encoding='utf-8')
        (public_dir / filename).write_text(json.dumps(public_community_report(report), indent=2, allow_nan=False) + '\n', encoding='utf-8')


class ReportError(Exception):
    """An actionable input problem, rather than a programming error."""


class OriginalPriceError(ReportError):
    """A failed original-price audit, also saved as a warning report."""


def property_address(row):
    parts = [str(row[k]).strip() for k in ("street_number", "street_name", "street_suffix")
             if row.get(k) is not None and str(row[k]).strip()]
    address = " ".join(parts) or "Address not supplied"
    if row.get("unit") is not None and str(row["unit"]).strip():
        address += " Unit " + str(row["unit"]).strip()
    return address + ", " + str(row.get("city", ""))


def normalize(value):
    return " ".join(str(value or "").split()).casefold()


def year_before(day):
    return day.replace(year=day.year - 1,
                       day=min(day.day, calendar.monthrange(day.year - 1, day.month)[1]))


def number(value, field, location, positive=False):
    if value is None or isinstance(value, bool):
        raise ReportError(f"{location}: missing or invalid {field}.")
    try:
        result = Decimal(str(value).replace(",", "").replace("$", "").strip())
    except InvalidOperation:
        raise ReportError(f"{location}: invalid {field}: {value!r}.") from None
    if not result.is_finite() or result < 0 or (positive and result == 0):
        raise ReportError(f"{location}: {field} must be {'positive' if positive else 'nonnegative'}.")
    return result


def closing_date(value, location):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ReportError(f"{location}: missing or unrecognized Closed Date: {value!r}.")


def file_hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_rows(path, columns):
    """Require an unambiguous header and a single data sheet; never save inputs."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        candidates = []
        required = set(columns.values())
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if row_number > 20:
                    break
                headers = [str(v).strip() if v is not None else "" for v in row]
                if required.issubset(headers):
                    candidates.append((sheet, row_number, headers))
                    break
        if len(candidates) != 1:
            raise ReportError(f"{path.name}: expected one data sheet with these headers: "
                              + ", ".join(sorted(required))
                              + ". Check the export columns and settings.json.")
        sheet, header_row, headers = candidates[0]
        if any(headers.count(name) != 1 for name in required):
            raise ReportError(f"{path.name}: duplicate required column headers.")
        result = []
        for row_number, values in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            if all(v is None or v == "" for v in values):
                continue
            raw = dict(zip(headers, values))
            row = {key: raw[name] for key, name in columns.items()}
            row["source"] = path.name
            row["row"] = row_number
            row["location"] = f"{path.name}, row {row_number}"
            result.append(row)
        return result, {"file": path.name, "sheet": sheet.title,
                        "header_row": header_row, "rows": len(result), "headers": headers}
    finally:
        workbook.close()


def make_lookup(settings):
    lookup = {}
    slugs = set()
    for community in settings["communities"]:
        slug = community["slug"]
        if not slug or any(c not in "abcdefghijklmnopqrstuvwxyz0123456789-" for c in slug):
            raise ReportError("Community slugs must contain only lowercase letters, digits and hyphens.")
        if slug in slugs:
            raise ReportError(f"Duplicate community slug: {slug}")
        slugs.add(slug)
        for alias in community["subdivisions"]:
            key = normalize(alias), normalize(community["city"])
            if key in lookup:
                raise ReportError(f"Overlapping community match: {alias}, {community['city']}.")
            lookup[key] = slug
    return lookup


def classify(rows_by_role, settings, as_of):
    lookup = make_lookup(settings)
    groups = {key: {normalize(v) for v in values}
              for key, values in settings["statuses"].items()}
    keys = list(groups)
    if any(groups[a] & groups[b] for i, a in enumerate(keys) for b in keys[i + 1:]):
        raise ReportError("Status groups overlap in settings.json.")
    buckets = {c["slug"]: {"active": [], "pending": [], "closed": []}
               for c in settings["communities"]}
    review, seen = [], set()
    original_price_errors = []
    start = year_before(as_of)
    for role, rows in rows_by_role.items():
        allowed = ["closed"] if role == "closed" else ["active", "pending"]
        for row in rows:
            location = row["location"]
            mls = str(row["mls"] or "").strip()
            if not mls:
                raise ReportError(f"{location}: missing MLS number.")
            if mls in seen:
                raise ReportError(f"{location}: duplicate MLS number {mls}. Resolve duplicate exports/records first.")
            seen.add(mls)
            status = normalize(row["status"])
            matches = [group for group in allowed if status in groups[group]]
            if len(matches) != 1:
                raise ReportError(f"{location}: unexpected status {row['status']!r} in {role} export. "
                                  "Review the status definition before updating settings.json.")
            group = matches[0]
            if not normalize(row["community"]) or not normalize(row["city"]):
                raise ReportError(f"{location}: missing Subdivision or City; community cannot be determined.")
            slug = lookup.get((normalize(row["community"]), normalize(row["city"])))
            reason = "included"
            if slug is None:
                reason = "community/city not in configured matches"
            elif group == "closed":
                closed = closing_date(row["closed_date"], location)
                if not start <= closed < as_of:
                    reason = "closing date outside reporting window"
            if reason == "included":
                if group in ("active", "closed"):
                    row["dom_value"] = number(row["dom"], "DOM", location)
                    row["list_value"] = number(row["list_price"], "List Price", location, positive=True)
                if group == "closed":
                    row["sold_value"] = number(row["sold_price"], "Sold Pr", location, positive=True)
                    try:
                        row["original_list_value"] = number(
                            row.get("original_list_price"), ORIGINAL_LIST_FIELD, location, positive=True)
                    except ReportError as error:
                        original_price_errors.append(
                            f"MLS {mls}, {row['community']}, {property_address(row)}: {error} "
                            f"Value: {row.get('original_list_price')!r}")
                buckets[slug][group].append(row)
            review.append({"source": row["source"], "row": row["row"], "mls": mls,
                           "subdivision": row["community"], "city": row["city"],
                           "status": row["status"], "community_slug": slug or "",
                           "disposition": reason})
    if original_price_errors:
        raise OriginalPriceError(
            f"WARNING - REPORT NOT GENERATED: Unusable Orig List Pr in {len(original_price_errors)} qualifying closed sale(s):\n"
            + "\n".join(original_price_errors)
            + "\nNo percentages calculated. Correct these records and rerun; List Price is never substituted.")
    return buckets, review


def rounded(value, digits=2):
    return None if value is None else round(float(value), digits)


def latest_month_bounds(as_of):
    """The previous calendar month, including when the report date is mid-month."""
    end_exclusive = as_of.replace(day=1)
    start = (end_exclusive - timedelta(days=1)).replace(day=1)
    return start, end_exclusive


def latest_month_statistics(closed, as_of):
    start, end = latest_month_bounds(as_of)
    monthly = [r for r in closed if start <= closing_date(r['closed_date'], r['location']) < end]
    total_sold = sum((r['sold_value'] for r in monthly), Decimal(0))
    total_original = sum((r['original_list_value'] for r in monthly), Decimal(0))
    return {
        'latest_complete_month': start.strftime('%Y-%m'),
        'latest_month_closed_sales': len(monthly),
        'latest_month_median_sold_price': rounded(statistics.median([r['sold_value'] for r in monthly])) if monthly else None,
        'latest_month_sale_to_list_percentage': rounded(total_sold / total_original * 100, 4) if monthly else None,
        'latest_month_median_closed_dom': rounded(statistics.median([r['dom_value'] for r in monthly])) if monthly else None,
    }


def monthly_history_months(as_of):
    if as_of.day != 1:
        raise ReportError("Monthly report date must be the first day of the month (YYYY-MM-01).")
    end = as_of.year * 12 + as_of.month - 1
    return [date(index // 12, index % 12 + 1, 1) for index in range(end - 12, end)]


def monthly_history_statistics(closed, as_of):
    months = monthly_history_months(as_of)
    for row in closed:
        # Validate every qualifying denominator; never fall back to List Price.
        number(row.get('original_list_value'), ORIGINAL_LIST_FIELD, row['location'], positive=True)
    history = []
    for start in months:
        end = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
        values = latest_month_statistics(closed, end)
        history.append(dict(zip(MONTHLY_HISTORY_FIELDS, (
            values['latest_complete_month'], values['latest_month_closed_sales'],
            values['latest_month_median_sold_price'], values['latest_month_sale_to_list_percentage'],
            values['latest_month_median_closed_dom'],
        ))))
    return history


def validate_monthly_history(report):
    """Validate public structure and reconciliation before local or GitHub publication."""
    try:
        day = date.fromisoformat(report['report_date'])
        months = [d.strftime('%Y-%m') for d in monthly_history_months(day)]
        history = report['monthly_history']
        if (report['closed_period_start'] != year_before(day).isoformat()
                or report['closed_period_end'] != (day - timedelta(days=1)).isoformat()
                or report['latest_complete_month'] != months[-1]):
            raise ReportError('Monthly history reporting period mismatch.')
        if type(history) is not list or len(history) != 12:
            raise ReportError('Monthly history must contain exactly 12 months.')
        for expected, row in zip(months, history):
            if type(row) is not dict or set(row) != set(MONTHLY_HISTORY_FIELDS):
                raise ReportError('Monthly history objects must contain exactly the five approved fields.')
            if row['month'] != expected:
                raise ReportError('Monthly history months must be consecutive and chronological.')
            count = row['closed_sales']
            if type(count) is not int or count < 0:
                raise ReportError('Monthly history closed sales must be a nonnegative integer.')
            for field in MONTHLY_HISTORY_FIELDS[2:]:
                value = row[field]
                if count == 0:
                    if value is not None:
                        raise ReportError('Zero-sales monthly statistics must be null.')
                elif (type(value) not in (int, float) or not math.isfinite(value)
                      or value < 0 or (field != 'median_closed_dom' and value == 0)):
                    raise ReportError('Monthly history contains an invalid statistic.')
        if sum(row['closed_sales'] for row in history) != report['closed_sales_12_months']:
            raise ReportError('Monthly closed-sales sum must equal the 12-month sales count.')
        latest = ('latest_complete_month', 'latest_month_closed_sales', 'latest_month_median_sold_price',
                  'latest_month_sale_to_list_percentage', 'latest_month_median_closed_dom')
        if any(history[-1][field] != report[key] for field, key in zip(MONTHLY_HISTORY_FIELDS, latest)):
            raise ReportError('Final monthly history record must match the latest-complete-month fields.')
    except (KeyError, TypeError, ValueError) as error:
        raise ReportError('Missing or invalid monthly history fields.') from error


def validate_source_dates(report):
    for key in ('inventory_updated', 'closed_sales_through'):
        value = report.get(key)
        try:
            if type(value) is not str or date.fromisoformat(value).isoformat() != value:
                raise ValueError()
        except (ValueError, TypeError):
            raise ReportError(f'{key} must be a valid YYYY-MM-DD date.') from None
    if report['closed_sales_through'] != report['closed_period_end']:
        raise ReportError('closed_sales_through must equal the closed reporting period end.')


def summarize(community, bucket, as_of, inventory_date=None):
    active, pending, closed = (bucket[k] for k in ("active", "pending", "closed"))
    total_sold = sum((r["sold_value"] for r in closed), Decimal(0))
    total_list = sum((r["list_value"] for r in closed), Decimal(0))
    total_original_list = sum((r["original_list_value"] for r in closed), Decimal(0))
    median = lambda values: statistics.median(values) if values else None
    result = {
        "community": community["name"], "city": community["city"], "state": "Illinois",
        "report_date": as_of.isoformat(), "closed_period_start": year_before(as_of).isoformat(),
        "closed_period_end": (as_of - timedelta(days=1)).isoformat(),
        "active_listings": len(active),
        "median_active_list_price": rounded(median([r["list_value"] for r in active])),
        "average_active_dom": rounded(statistics.mean([r["dom_value"] for r in active])) if active else None,
        "pending_contingent_fin_listings": len(pending),
        "closed_sales_12_months": len(closed),
        "median_sold_price": rounded(median([r["sold_value"] for r in closed])),
        "aggregate_sale_to_list_percentage": rounded(total_sold / total_original_list * 100, 4) if total_original_list else None,
        "median_closed_dom": rounded(median([r["dom_value"] for r in closed])),
        "sale_to_list_basis": SALE_TO_LIST_DEFINITION,
        "sale_to_list_denominator_field": ORIGINAL_LIST_FIELD,
        "total_closed_sold_price": rounded(total_sold),
        "total_closed_list_price": rounded(total_list),
        "total_closed_original_list_price": rounded(total_original_list),
        "included_mls_numbers": {k: [str(r["mls"]) for r in rows] for k, rows in bucket.items()},
        **latest_month_statistics(closed, as_of),
        "monthly_history": monthly_history_statistics(closed, as_of),
        "inventory_updated": (inventory_date or as_of).isoformat(),
        "closed_sales_through": (as_of - timedelta(days=1)).isoformat(),
    }
    validate_monthly_history(result)
    return result


def write_excel(path, reports, metadata):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Market Summary"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet["A2"] = "55+ Community Market Report"
    sheet["A2"].font = Font(name="Arial", size=16, bold=True, color="243746")
    sheet["A3"] = f"Report date: {metadata['report_date']} | Closed sales: {metadata['closed_period_start']} through {metadata['closed_period_end']}"
    headers = ["Community", "City", "Active listings", "Median active list price",
               "Average active DOM", "Pending / contingent / FIN", "Closed sales (12 months)",
               "Median sold price", "Sale-to-original-list %", "Median closed DOM"]
    for column, value in enumerate(headers, 1):
        sheet.cell(5, column, value)
    fields = ["community", "city", "active_listings", "median_active_list_price",
              "average_active_dom", "pending_contingent_fin_listings", "closed_sales_12_months",
              "median_sold_price", "aggregate_sale_to_list_percentage", "median_closed_dom"]
    for row_number, report in enumerate(reports, 6):
        for column, field in enumerate(fields, 1):
            value = report[field]
            if field == "aggregate_sale_to_list_percentage" and value is not None:
                value /= 100
            sheet.cell(row_number, column, "n.a." if value is None else value)
    end = 5 + len(reports)
    table = Table(displayName="CommunitySummary", ref=f"A5:J{end}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    widths = [24, 18, 13, 19, 17, 22, 18, 18, 20, 17]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=5, max_row=end, max_col=10):
        for cell in row:
            cell.font = Font(name="Arial", size=11, color="243746")
            cell.alignment = Alignment(vertical="center", horizontal="left" if cell.column <= 2 else "right")
    for cell in sheet[5]:
        cell.font = Font(name="Arial", size=11, color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="243746")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[5].height = 45
    for row in range(6, end + 1):
        sheet.row_dimensions[row].height = 27
        for column in (4, 8):
            sheet.cell(row, column).number_format = '"$"#,##0'
        for column in (5, 10):
            sheet.cell(row, column).number_format = "0.0"
        sheet.cell(row, 9).number_format = "0.00%"
    notes = [
        "DOM uses MT. Pending/contingent/FIN combines the configured status codes.",
        "Sale-to-list %: " + SALE_TO_LIST_DEFINITION,
        "n.a. means no qualifying listings. Counts of zero mean no matching records in the supplied exports.",
        "Active/pending figures reflect the export snapshot; the report date only filters closed sales.",
        "Sources: " + "; ".join(s["file"] for s in metadata["sources"]),
        f"{metadata['excluded_rows']} source rows excluded. See listing_review.csv for every row's disposition.",
    ]
    for offset, note in enumerate(notes, end + 3):
        sheet.cell(offset, 1, note).font = Font(name="Arial", size=10, color="52616B")
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = f"A1:J{end + 8}"
    add_definitions(workbook)
    workbook.save(path)
    workbook.close()


def add_definitions(workbook):
    definitions = workbook.create_sheet("Definitions")
    definitions.sheet_view.showGridLines = False
    definitions["A2"] = "Metric definitions"
    definitions["A2"].font = Font(name="Arial", size=16, bold=True, color="243746")
    definitions.append([])
    entries = [
        ("Metric", "Definition"),
        ("Active listings", "Count of qualifying listings with ACTV, NEW or PCHG status."),
        ("Median active list price", "Median List Price of qualifying active listings."),
        ("Average active DOM", "Arithmetic average of MT for qualifying active listings."),
        ("Pending / contingent / FIN", "Combined count of A/I, PEND, HC24, HC48, HS48, FIN and SS listings."),
        ("Closed sales (12 months)", "Count of qualifying CLSD listings in the preceding 12 months, excluding the report date."),
        ("Median sold price", "Median Sold Pr for qualifying closed sales."),
        ("Sale-to-List %", SALE_TO_LIST_DEFINITION),
        ("Median closed DOM", "Median MT for qualifying closed sales."),
        ("Original list price field", "Orig List Pr in the closed export. Every qualifying value must be positive and numeric. No List Price fallback."),
        ("Community matching", "Exact configured subdivision and city, ignoring capitalization and extra spaces."),
        ("Validation warnings", "Every run checks qualifying original prices. Unusable values produce a warning report with MLS numbers and addresses; market reports are withheld."),
        ("Percentage display", "Display Sale-to-List % to two decimal places. Calculate the aggregate from source amounts before rounding."),
        ("Approved methodology", "Use Orig List Pr permanently. Change this methodology only on the user's explicit instruction."),
        ("Report date", "The report's as-of date. Closed sales include the preceding 12 calendar months and exclude the report date; active and pending reflect the export snapshot."),
        ("Community and city", "Configured community name and city. Shorewood Glen matches Shorewood Glen Del Webb in Shorewood; other names match their configured subdivision and city."),
        ("Unavailable statistics", "No qualifying observations: counts are zero; price, DOM and percentage statistics are unavailable (blank in history, n.a. in the summary)."),
        ("Latest complete month (JSON)", "The complete calendar month before the report date, labeled YYYY-MM. Its closed count, median Sold Pr, aggregate Sold Pr / Orig List Pr percentage and median MT are separate from the 12-month metrics."),
        ("Latest month with no sales", "Count is zero; monthly median sold price, sale-to-list percentage and median closed DOM are null. No previous values are carried forward."),
        ("Inventory updated (JSON)", "Inventory source refresh date, YYYY-MM-DD. Daily updates change this date only with the inventory metrics. Monthly runs preserve the known date for unchanged inventory sources or use the supplied refresh date / processing date."),
        ("Closed sales through (JSON)", "Final date included in the closed reporting window, YYYY-MM-DD; equals closed_period_end. This is the coverage cutoff, not the date of the last individual sale. Daily inventory updates preserve it."),
    ]
    for row_number, entry in enumerate(entries, 5):
        for column, value in enumerate(entry, 1):
            cell = definitions.cell(row_number, column, value)
            cell.font = Font(name="Arial", size=11, color="243746")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        definitions.row_dimensions[row_number].height = 44
    for cell in definitions[5]:
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
    definitions.column_dimensions["A"].width = 30
    definitions.column_dimensions["B"].width = 100


HISTORY_COLUMNS = [
    ("Report Date", "report_date"), ("Community", "community"), ("City", "city"),
    ("Active Listings", "active_listings"), ("Median Active List Price", "median_active_list_price"),
    ("Average Active DOM", "average_active_dom"), ("Pending / Contingent / FIN", "pending_contingent_fin_listings"),
    ("Closed Sales (12 Months)", "closed_sales_12_months"), ("Median Sold Price", "median_sold_price"),
    ("Sale-to-List %", "aggregate_sale_to_list_percentage"), ("Median Closed DOM", "median_closed_dom"),
]


def prepare_history(history_path, staged_path, reports, as_of):
    """Append to a staged copy only; never assign to existing historical rows."""
    expected_hash = file_hash(history_path) if history_path.exists() else None
    wb = openpyxl.load_workbook(history_path) if expected_hash else openpyxl.Workbook()
    try:
        if expected_hash:
            if wb.sheetnames != ["Market History", "Definitions"]:
                raise ReportError("Unexpected market_history.xlsx worksheets; history was not changed.")
            sheet = wb["Market History"]
            if [c.value for c in sheet[1]] != [c[0] for c in HISTORY_COLUMNS]:
                raise ReportError("Unexpected history columns; history was not changed.")
        else:
            sheet = wb.active
            sheet.title = "Market History"
            sheet.append([c[0] for c in HISTORY_COLUMNS])
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = "D2"
            widths = [15, 24, 18, 13, 19, 17, 23, 20, 18, 18, 18]
            for col, width in enumerate(widths, 1):
                sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
            for cell in sheet[1]:
                cell.font = Font(name="Arial", size=11, color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="243746")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.row_dimensions[1].height = 45
            add_definitions(wb)
        seen = set()
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in values):
                raise ReportError("Unexpected blank row in history; history was not changed.")
            day = closing_date(values[0], "Market History")
            key = day, normalize(values[1])
            if key in seen:
                raise ReportError("Duplicate Report Date + Community already in history; no changes made.")
            seen.add(key)
            if day == as_of:
                raise ReportError(f"WARNING: Report date {as_of} already exists in market_history.xlsx. No historical rows changed.")
        incoming = set()
        for report in reports:
            key = as_of, normalize(report['community'])
            if report['report_date'] != as_of.isoformat() or key in seen or key in incoming:
                raise ReportError("Duplicate community or inconsistent report date in incoming history records.")
            if report.get('sale_to_list_denominator_field') != ORIGINAL_LIST_FIELD:
                raise ReportError("History requires the approved Orig List Pr methodology.")
            incoming.add(key)
            values = [report[field] for _, field in HISTORY_COLUMNS]
            values[0] = as_of
            values[9] = values[9] / 100 if values[9] is not None else None
            sheet.append(values)
            row = sheet.max_row
            for cell in sheet[row]:
                cell.font = Font(name="Arial", size=11, color="243746")
                cell.alignment = Alignment(vertical="center", horizontal="left" if cell.column <= 3 else "right")
            sheet.row_dimensions[row].height = 27
            sheet.cell(row, 1).number_format = 'yyyy-mm-dd'
            for col in (5, 9):
                sheet.cell(row, col).number_format = '"$"#,##0'
            for col in (6, 11):
                sheet.cell(row, col).number_format = '0.0'
            sheet.cell(row, 10).number_format = '0.00%'
        ref = f"A1:K{sheet.max_row}"
        if 'MarketHistory' in sheet.tables:
            sheet.tables['MarketHistory'].ref = ref
            sheet.tables['MarketHistory'].autoFilter.ref = ref
        else:
            table = Table(displayName='MarketHistory', ref=ref)
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            sheet.add_table(table)
        wb.save(staged_path)
        return expected_hash
    finally:
        wb.close()


def commit_history(staged_path, history_path, expected_hash):
    actual = file_hash(history_path) if history_path.exists() else None
    if actual != expected_hash:
        raise ReportError("History changed during this run. No history update was committed; retry after reviewing it.")
    staged_path.replace(history_path)


def import_archive_history(archive_dir):
    """Import an already validated archive without regenerating any market reports."""
    metadata = json.loads((archive_dir / 'run_info.json').read_text())
    audit = metadata['original_list_price_validation']
    if audit['field'] != ORIGINAL_LIST_FIELD or audit['unusable_records'] or audit['fallback_used']:
        raise ReportError("Archive does not show a successful original-price validation.")
    report_folder = archive_dir / 'audit' / 'communities'
    if not report_folder.exists():
        report_folder = archive_dir / 'communities'  # Original private archives remain readable.
    reports = [json.loads((report_folder / (c['slug'] + '.json')).read_text())
               for c in metadata['settings']['communities']]
    if sum(r['closed_sales_12_months'] for r in reports) != audit['qualifying_closed_sales']:
        raise ReportError("Archive closed counts do not match the validation record.")
    history = archive_dir.parent / 'market_history.xlsx'
    with tempfile.TemporaryDirectory(prefix='.history-', dir=archive_dir.parent) as temporary:
        stage = Path(temporary) / 'market_history.xlsx'
        expected = prepare_history(history, stage, reports, date.fromisoformat(metadata['report_date']))
        commit_history(stage, history, expected)
    return history, reports


def publish_reports(stage, archive_stage, output_dir, archive_dir, history_stage=None, history_hash=None):
    """Publish complete folders, preserving the previous Output and rolling back on failure."""
    if archive_dir.exists():
        raise ReportError(f"WARNING: Archive already exists: {archive_dir}. Nothing was overwritten. "
                          "Review the existing archive before choosing a different report date.")
    previous = None
    archived = False
    published = False
    try:
        if output_dir.exists():
            previous_root = archive_dir.parent / "_previous_output"
            previous_root.mkdir(exist_ok=True)
            previous = previous_root / f"preserved_{datetime.now():%Y%m%d_%H%M%S_%f}"
            # Both paths were checked by run(); never move an unchecked computed root.
            output_dir.rename(previous)
        archive_stage.rename(archive_dir)
        archived = True
        stage.rename(output_dir)
        published = True
        if history_stage is not None:
            commit_history(history_stage, archive_dir.parent / 'market_history.xlsx', history_hash)
    except (OSError, ReportError):
        if published:
            output_dir.rename(stage)
        if archived:
            archive_dir.rename(archive_stage)
        if previous is not None and previous.exists():
            previous.rename(output_dir)
        raise


def run(input_dir, output_dir, settings_path, as_of, archive_root=None, inventory_date=None):
    monthly_history_months(as_of)  # Reject partial-month reporting before any writes.
    input_dir, output_dir = input_dir.resolve(), output_dir.absolute()
    archive_root = (archive_root or output_dir.parent / "Archive").resolve()
    if output_dir.is_symlink() or (hasattr(output_dir, "is_junction") and output_dir.is_junction()):
        raise ReportError("Output must be a regular folder, not a link or junction.")
    output_dir = output_dir.resolve()
    roots = [input_dir, output_dir, archive_root]
    if any(a == b or a in b.parents or b in a.parents
           for i, a in enumerate(roots) for b in roots[i + 1:]):
        raise ReportError("Input, Output and Archive must be separate folders, with none inside another.")
    archive_dir = archive_root / as_of.isoformat()
    if archive_dir.exists():
        raise ReportError(f"WARNING: Archive already exists: {archive_dir}. "
                          "The run was stopped; current and historical reports were not overwritten.")
    settings = json.loads(settings_path.read_text(encoding="utf-8-sig"))
    if not settings.get("rules_confirmed", False):
        raise ReportError("Reporting rules are not confirmed. Review settings.json before running.")
    if settings.get("closed_columns", {}).get("original_list_price") != ORIGINAL_LIST_FIELD:
        raise ReportError("Approved methodology requires Orig List Pr. A different denominator mapping is not permitted.")
    if input_dir.resolve() == output_dir.resolve() or input_dir.resolve() in output_dir.resolve().parents:
        raise ReportError("Output must be outside the Input folder.")
    paths = {role: input_dir / name for role, name in settings["input_files"].items()}
    if set(paths) != {"inventory", "closed"} or len(set(p.resolve() for p in paths.values())) != 2:
        raise ReportError("Configure two different input files: inventory and closed.")
    for path in paths.values():
        if path.parent.resolve() != input_dir.resolve() or path.suffix.lower() != ".xlsx":
            raise ReportError("Input file names must be .xlsx files directly inside Input.")
    existing = {p.name for p in input_dir.glob("*.xlsx") if not p.name.startswith("~$")}
    if existing != {p.name for p in paths.values()}:
        raise ReportError("Input must contain exactly the two .xlsx files named in settings.json. "
                          "Move old exports to Archive and use the expected file names.")
    rows_by_role, sources = {}, []
    for role, path in paths.items():
        before = file_hash(path)
        columns = dict(settings["columns"])
        if role == "closed":
            columns["original_list_price"] = ORIGINAL_LIST_FIELD
        rows, source = load_rows(path, columns)
        source.update(role=role, sha256=before)
        rows_by_role[role] = rows
        sources.append(source)
    try:
        buckets, review = classify(rows_by_role, settings, as_of)
    except OriginalPriceError as error:
        warning_dir = archive_root / "_validation_warnings"
        warning_dir.mkdir(parents=True, exist_ok=True)
        warning_path = warning_dir / f"{as_of.isoformat()}_WARNING_{datetime.now():%Y%m%d_%H%M%S_%f}.txt"
        warning_path.write_text(
            f"Report date: {as_of.isoformat()}\n{SALE_TO_LIST_DEFINITION}\n\n{error}\n",
            encoding="utf-8")
        raise OriginalPriceError(f"{error}\nWarning details saved to: {warning_path}") from None
    # Inventory provenance is independent of the monthly closed-sales cutoff.
    # Reusing the same monthly source preserves its known refresh date.
    if inventory_date is None:
        inventory_date = date.today()
        previous_info = output_dir / 'run_info.json'
        if previous_info.is_file():
            previous = json.loads(previous_info.read_text(encoding='utf-8'))
            prior_hashes = {s['sha256'] for s in previous.get('sources', []) if s.get('role') == 'inventory'}
            current_hash = next(s['sha256'] for s in sources if s['role'] == 'inventory')
            if current_hash in prior_hashes:
                dates = {json.loads((output_dir / 'communities' / (c['slug'] + '.json')).read_text(encoding='utf-8')).get('inventory_updated')
                         for c in settings['communities']}
                if len(dates) == 1 and None not in dates:
                    inventory_date = date.fromisoformat(dates.pop())
    reports = [summarize(c, buckets[c["slug"]], as_of, inventory_date) for c in settings["communities"]]
    warnings = []
    for c in settings["communities"]:
        for role in ("inventory", "closed"):
            if not any(normalize(r["community"]) in {normalize(a) for a in c["subdivisions"]}
                       and normalize(r["city"]) == normalize(c["city"]) for r in rows_by_role[role]):
                warnings.append(f"{c['name']}: no matching records in {role} export. Verify search coverage.")
    metadata = {
        "report_date": as_of.isoformat(), "closed_period_start": year_before(as_of).isoformat(),
        "closed_period_end": (as_of - timedelta(days=1)).isoformat(),
        "generated_at": datetime.now().astimezone().isoformat(), "sources": sources,
        "archive_folder": str(archive_dir),
        "communities_processed": len(reports),
        "settings": settings, "warnings": warnings,
        "original_list_price_validation": {
            "field": settings["closed_columns"]["original_list_price"],
            "qualifying_closed_sales": sum(len(b["closed"]) for b in buckets.values()),
            "unusable_records": [], "fallback_used": False,
        },
        "excluded_rows": sum(r["disposition"] != "included" for r in review),
        "status_counts": {role: dict(Counter(str(r['status']) for r in rows)) for role, rows in rows_by_role.items()},
        "coverage_note": "Export search coverage and snapshot date must be verified by the user; file rows cannot prove completeness.",
    }
    for source in sources:
        if file_hash(input_dir / source["file"]) != source["sha256"]:
            raise ReportError("An input file changed while being read. Please rerun with stable exports.")
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    archive_root.mkdir(parents=True, exist_ok=True)
    # Stage both complete folders before touching current or historical reports.
    with tempfile.TemporaryDirectory(prefix=".building-", dir=output_dir.parent) as temporary:
        stage = Path(temporary) / 'reports'
        stage.mkdir()
        for community, report in zip(settings["communities"], reports):
            report["warnings"] = [w for w in warnings if w.startswith(community["name"] + ":")]
        write_community_reports(stage, settings['communities'], reports)
        (stage / "run_info.json").write_text(json.dumps(metadata, indent=2) + "\n", encoding="utf-8")
        with (stage / "listing_review.csv").open("w", newline="", encoding="utf-8-sig") as handle:
            fields = ["source", "row", "mls", "subdivision", "city", "status", "community_slug", "disposition"]
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            # Escape spreadsheet formula prefixes in untrusted source text.
            writer.writerows({k: "'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v
                              for k, v in row.items()} for row in review)
        write_excel(stage / "market_summary.xlsx", reports, metadata)
        history_stage = Path(temporary) / 'market_history.xlsx'
        history_hash = prepare_history(archive_root / 'market_history.xlsx', history_stage, reports, as_of)
        with tempfile.TemporaryDirectory(prefix=".archiving-", dir=archive_root) as archive_temporary:
            archive_stage = Path(archive_temporary)
            shutil.copytree(stage, archive_stage, dirs_exist_ok=True)
            source_folder = archive_stage / "source_exports"
            source_folder.mkdir()
            for source in sources:
                copied = source_folder / source["file"]
                shutil.copy2(input_dir / source["file"], copied)
                if file_hash(copied) != source["sha256"] or file_hash(input_dir / source["file"]) != source["sha256"]:
                    raise ReportError("An input changed before archival. No reports published; rerun with stable exports.")
            publish_reports(stage, archive_stage, output_dir, archive_dir, history_stage, history_hash)
    return output_dir, metadata


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--as-of", type=date.fromisoformat, default=date.today(), metavar="YYYY-MM-DD",
                        help="Report date; sales from the prior 12 months, excluding this date (default: today).")
    parser.add_argument("--input", type=Path, default=BASE / "Input")
    parser.add_argument("--output", type=Path, default=BASE / "Output")
    parser.add_argument("--settings", type=Path, default=BASE / "settings.json")
    parser.add_argument("--inventory-date", type=date.fromisoformat,
                        help="Inventory source refresh date; otherwise preserve the known date for unchanged sources or use today's update date.")
    parser.add_argument("--archive", type=Path, default=None,
                        help="Archive folder (default: Archive beside Output). Existing report dates are never overwritten.")
    args = parser.parse_args()
    try:
        final, metadata = run(args.input, args.output, args.settings, args.as_of, args.archive,
                              inventory_date=args.inventory_date)
    except (ReportError, OSError, ValueError, KeyError) as error:
        print(f"REPORT NOT CREATED: {error}\nOriginal exports and previous reports were not changed.", file=sys.stderr)
        print("Monthly report: FAILED\nArchive: NOT CONFIRMED\nHistory: NOT CONFIRMED\nJSON validation: NOT ATTEMPTED", file=sys.stderr)
        print("GitHub publication: NOT ATTEMPTED (monthly run did not complete).", file=sys.stderr)
        return 1
    print("Monthly report: SUCCESSFUL")
    print("Archive: SUCCESSFUL")
    print("History: SUCCESSFUL")
    print(f"Report date: {metadata['report_date']}")
    print(f"Communities processed: {metadata['communities_processed']}")
    print(f"Closed sales analyzed: {metadata['original_list_price_validation']['qualifying_closed_sales']}")
    print(f"Current reports: {final}")
    print(f"Archive folder created: {metadata['archive_folder']}")
    print(f"History updated: {Path(metadata['archive_folder']).parent / 'market_history.xlsx'}")
    print("Data validation warnings: " + (f"Yes ({len(metadata['warnings'])})" if metadata['warnings'] else "None"))
    for warning in metadata["warnings"]:
        print("CHECK: " + warning)
    from github_publisher import after_monthly
    publication = after_monthly(BASE, final, metadata)
    print("JSON validation: " + publication['json_validation'].upper())
    print("GitHub publication: " + publication['github_publication'].upper())
    print(publication['message'])
    if publication.get('log_warning'):
        print(publication['log_warning'])
    if publication['github_publication'] == 'failed':
        return 2  # Local reports succeeded. Publication failure cannot undo them.
    return 0


if __name__ == "__main__":
    sys.exit(main())
